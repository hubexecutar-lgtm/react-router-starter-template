#!/usr/bin/env python3
"""Testes dos scripts determinísticos (extração e validação) sobre a fixture FICTÍCIA.

Execução: python3 tests/test_scripts.py   (ou pytest, se disponível)
"""
from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
PY = sys.executable
SCHEMA = ["dependency_id", "source_artifact_id", "target_artifact_id", "relation", "mandatory",
          "gate_id", "required_status", "status"]

# Registro de referência para a fixture fictícia (não representa o EXECUTAR HUB real).
REGISTRO = [
    ["DEP-001", "A00", "A01", "produz_insumo", "TRUE", "G-01", "concluido", "ativo"],
    ["DEP-002", "A01", "A02", "produz_insumo", "TRUE", "G-01", "concluido", "ativo"],
    ["DEP-003", "A02", "A03", "produz_insumo", "TRUE", "G-01", "concluido", "ativo"],
    ["DEP-004", "A03", "A04", "produz_insumo", "TRUE", "G-02", "concluido", "ativo"],
    ["DEP-005", "A04", "A03", "produz_insumo", "FALSE", "G-02", "concluido", "ativo"],
    ["DEP-006", "A04", "A05", "produz_insumo", "TRUE", "G-03", "concluido", "ativo"],
    ["DEP-007", "A02", "A05", "produz_insumo", "TRUE", "G-03", "concluido", "ativo"],
    ["DEP-008", "A00", "A06", "produz_insumo", "TRUE", "G-04", "concluido", "ativo"],
]
ANEXO = [
    ["DEP-001", "DIRECT", "16_REG_Dependencias linha DEP-001"],
    ["DEP-002", "DERIVED", "15_REG_Campos: C-A02-01 consome C-A01-01"],
    ["DEP-003", "DIRECT", "19_REG_Documentos DOC-01"],
    ["DEP-004", "DERIVED", "15_REG_Campos: C-A04-01 consome C-A03-01"],
    ["DEP-005", "DERIVED", "15_REG_Campos: C-A03-02 consome C-A04-02 (informativa; quebra o ciclo A03↔A04)"],
    ["DEP-006", "DIRECT", "20_REG_Decisoes DEC-01"],
    ["DEP-007", "DERIVED", "15_REG_Campos: C-A05-01 consome C-A02-02"],
    ["DEP-008", "DERIVED", "15_REG_Campos: C-A06-02 consome C-A00-02"],
    ["DEP-009", "CONFLICT", "20_REG_Decisoes DEC-02 × 19_REG_Documentos DOC-02 (A05 → A06)"],
]
FASES = [{"fase": "F1", "artefatos": ["A00", "A01", "A02"]}, {"fase": "F2", "artefatos": ["A03"]},
         {"fase": "F3", "artefatos": ["A04", "A05"]}, {"fase": "F4", "artefatos": ["A06"]}]


def rodar(*args: object) -> subprocess.CompletedProcess:
    return subprocess.run([PY, *map(str, args)], capture_output=True, text=True)


def preparar(td: Path, registro=None, anexo=None, fases=None, shared=False, cabecalho=None) -> dict[str, Path]:
    xlsx = td / "cp.xlsx"
    r = rodar(RAIZ / "tests/fixtures/gerar_control_plane_ficticio.py", xlsx, *(["--shared"] if shared else []))
    assert r.returncode == 0, r.stderr
    cp = td / "cp.json"
    r = rodar(RAIZ / "scripts/extrair_control_plane.py", xlsx, cp)
    assert r.returncode == 0, r.stderr
    reg, anx, fz = td / "registro.csv", td / "anexo.csv", td / "fases.json"
    with reg.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(SCHEMA if cabecalho is None else cabecalho)
        w.writerows(REGISTRO if registro is None else registro)
    with anx.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["dependency_id", "status_epistemico", "justificativa"])
        w.writerows(ANEXO if anexo is None else anexo)
    fz.write_text(json.dumps(FASES if fases is None else fases), encoding="utf-8")
    return {"xlsx": xlsx, "cp": cp, "reg": reg, "anx": anx, "fases": fz}


def validar(p: dict[str, Path], com_fases=True) -> subprocess.CompletedProcess:
    extra = ["--fases", p["fases"]] if com_fases else []
    return rodar(RAIZ / "scripts/validar_registro.py", "--registro", p["reg"], "--anexo", p["anx"],
                 "--control-plane", p["cp"], *extra)


def test_extracao_inline_e_shared_equivalentes():
    with tempfile.TemporaryDirectory() as a, tempfile.TemporaryDirectory() as b:
        pa, pb = preparar(Path(a)), preparar(Path(b), shared=True)
        ja, jb = json.loads(pa["cp"].read_text()), json.loads(pb["cp"].read_text())
        assert ja["abas"] == jb["abas"]
        aus = ja["obrigatorias"]["ausentes"]
        assert "12_REG_Portfolio" in aus and "01_Formulario_A12" in aus and "16_REG_Dependencias" not in aus
        assert ja["abas"]["15_REG_Campos"]["total_registros"] == 14


def test_registro_de_referencia_valido():
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td)))
        assert r.returncode == 0, r.stdout
        assert "cobertura de campos: 14/14" in r.stdout
        assert "camada 1: A00" in r.stdout
        assert "DEP-009: CONFLICT" not in r.stdout  # CONFLICT fica só no anexo, fora do registro
        assert "anexo: DEP-009 não está no registro" in r.stdout


def test_ciclo_bloqueante_e_erro():
    reg = [l[:] for l in REGISTRO]
    reg[4][4] = "TRUE"  # DEP-005 A04→A03 vira bloqueante → ciclo A03↔A04
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), registro=reg), com_fases=False)
        assert r.returncode == 1 and "ciclo bloqueante" in r.stdout, r.stdout


def test_derived_sem_justificativa_e_erro():
    anx = [l[:] for l in ANEXO]
    anx[1][2] = ""
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), anexo=anx))
        assert r.returncode == 1 and "DEP-002: DERIVED sem justificativa" in r.stdout, r.stdout


def test_gate_inexistente_e_erro():
    reg = [l[:] for l in REGISTRO]
    reg[1][5] = "G-99"
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), registro=reg))
        assert r.returncode == 1 and "G-99" in r.stdout, r.stdout


def test_reuso_de_id_canonico_com_outro_par_e_erro():
    reg = [l[:] for l in REGISTRO]
    reg[0][1], reg[0][2] = "A01", "A02"  # DEP-001 existe em 16_REG como A00→A01
    reg = [reg[0]] + reg[2:]
    anx = [a for a in ANEXO if a[0] != "DEP-002"]
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), registro=reg, anexo=anx), com_fases=False)
        assert r.returncode == 1 and "preservar IDs canônicos" in r.stdout, r.stdout


def test_artefato_inexistente_e_erro():
    reg = [l[:] for l in REGISTRO] + [["DEP-010", "A00", "A99", "produz_insumo", "TRUE", "G-01", "concluido", "ativo"]]
    anx = ANEXO + [["DEP-010", "PROPOSED", "teste"]]
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), registro=reg, anexo=anx), com_fases=False)
        assert r.returncode == 1 and "A99" in r.stdout, r.stdout


def test_fase_antes_da_dependencia_bloqueante_e_erro():
    fases = [{"fase": "F1", "artefatos": ["A00", "A01", "A02", "A04"]}, {"fase": "F2", "artefatos": ["A03"]},
             {"fase": "F3", "artefatos": ["A05"]}, {"fase": "F4", "artefatos": ["A06"]}]
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), fases=fases))
        assert r.returncode == 1 and "vem antes de sua dependência bloqueante A03" in r.stdout, r.stdout


def test_fase_incompleta_e_erro():
    fases = FASES[:-1]
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), fases=fases))
        assert r.returncode == 1 and "artefatos sem fase: ['A06']" in r.stdout, r.stdout


def test_gap_no_registro_e_erro():
    anx = [l[:] for l in ANEXO]
    anx[7][1] = "GAP"
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), anexo=anx))
        assert r.returncode == 1 and "GAP não vira linha" in r.stdout, r.stdout


def test_cabecalho_fora_do_schema_e_erro():
    with tempfile.TemporaryDirectory() as td:
        r = validar(preparar(Path(td), cabecalho=SCHEMA[:-1] + ["estado"]), com_fases=False)
        assert r.returncode == 1 and "cabeçalho deve ser exatamente" in r.stdout, r.stdout


def test_extrator_confere_com_openpyxl_quando_disponivel():
    try:
        import openpyxl  # noqa: F401 — referência independente; o extrator usa só stdlib
    except ImportError:
        print("  (openpyxl ausente: teste de conformidade pulado)")
        return
    with tempfile.TemporaryDirectory() as td:
        p = preparar(Path(td), shared=True)
        wb = openpyxl.load_workbook(p["xlsx"])
        ext = json.loads(p["cp"].read_text())["abas"]
        assert sorted(wb.sheetnames) == sorted(ext)
        for nome in ("15_REG_Campos", "16_REG_Dependencias", "17_REG_Gates"):
            def aparar(l):  # openpyxl completa até a largura máxima; o extrator corta vazios à direita
                l = [str(x) for x in l]
                while l and l[-1] == "":
                    l.pop()
                return l
            linhas = [aparar("" if c is None else c for c in row) for row in wb[nome].iter_rows(values_only=True)]
            assert [aparar(l) for l in ext[nome]["linhas_brutas"]] == linhas, nome


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"ok     {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHOU {nome}\n{e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
